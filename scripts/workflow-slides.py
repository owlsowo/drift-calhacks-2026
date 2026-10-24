"""Build editable PowerPoint source for import through Figma Slides' normal UI.
No generated imagery: all pictures are captured project/browser evidence.
"""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT.parent / 'outputs' / 'drift-workflow-source.pptx'
INK = '193D43'
PAPER = 'F2EDD9'
MUTED = '4D6564'
GOLD = 'C1A96C'
APP = 'https://drift-game.drift-calhacks-2026-owlsowo.workers.dev'
FIGMA = 'https://www.figma.com/make/4Sk3BHNmaGuxH17B36jAxP/Drift-Bottle-Interaction-Prototype'
REPO = 'https://github.com/owlsowo/drift-calhacks-2026'
r = Presentation()
r.slide_width, r.slide_height = Inches(13.333), Inches(7.5)

def text(s, x, y, w, h, body, size=22, color=INK, serif=False, bold=False, link=None):
    shape = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(.01)
    tf.margin_top = tf.margin_bottom = Inches(.01)
    for i, line in enumerate(body.split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.name = 'Georgia' if serif else 'Arial'
        p.font.size, p.font.bold = Pt(size), bold
        p.font.color.rgb = RGBColor.from_string(color)
        p.space_after = Pt(12)
        if link:
            p.runs[0].hyperlink.address = link
    return shape

def slide(kicker, title, number):
    s = r.slides.add_slide(r.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = RGBColor.from_string(PAPER)
    text(s,.55,.35,12,.4,kicker.upper(),12,MUTED,bold=True)
    text(s,.55,.95,12,1.35,title,38,serif=True)
    text(s,.55,7.05,11,.25,'DRIFT / CAL HACKS 13.0 / ACTUAL BUILD & WORKFLOW',10,MUTED)
    text(s,12,7.02,.7,.3,f'{number:02}',12,MUTED)
    return s

def picture(s,path,x,y,w):
    return s.shapes.add_picture(str(ROOT/path), Inches(x), Inches(y), width=Inches(w))

def panel(s,x,y,w,h,title,body):
    shape=s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid();shape.fill.fore_color.rgb=RGBColor.from_string('E5DECA')
    shape.line.fill.background()
    text(s,x+.2,y+.18,w-.4,.5,title,22,bold=True)
    text(s,x+.2,y+.85,w-.4,h-1,body,19)

s=slide('Six words from you. Six from someone else.','A little story, together.',1)
text(s,.55,2.55,4.25,1.4,'Reviving the QQ / WeChat drift bottle as a daily creative ritual.',26,serif=True)
text(s,.55,4.3,4.05,1.4,'Catch a beginning.\nFinish it in six words.\nBoth people keep the postcard.',21)
picture(s,Path('docs/evidence/workflow/live-empty.png'),4.95,2.45,7.8)
text(s,5,6.93,7,.25,'Actual hosted app • no AI-generated imagery',10,MUTED)
s.notes_slide.notes_text_frame.text='Audience hypothesis: first-time event attendees taking a short creative break. No attendee interviews or measured social outcomes are claimed. Photo credits are linked in the app. Live app: '+APP

s=slide('A small shared activity','Human words. A shared souvenir.',2)
panel(s,.55,2.4,3.9,3.8,'01 / Catch','Shake the bottle or use the button. A server-held claim prevents two writers completing the same catch.')
panel(s,4.72,2.4,3.9,3.8,'02 / Add six words','The person writes the ending. Optional Claude guidance suggests an angle; it does not supply the finished reply.')
panel(s,8.89,2.4,3.9,3.8,'03 / Keep & pass on','The original writer receives the completed postcard live. Tossing a new beginning is optional after finishing.')
s.notes_slide.notes_text_frame.text='Demonstrate on two identities using fresh text. The hosted two-browser test used synthetic test writers with real Claude calls. A separate local recording is explicitly labeled sandbox; do not call it the hosted Claude demo.'

s=slide('Figma Make → implementation','An empty sea should still feel welcoming.',3)
picture(s,Path('docs/evidence/workflow/figma-empty.png'),.55,2.35,4.8)
text(s,5.8,2.45,6.8,.7,'The first prototype was too tall and opaque.',25,serif=True)
text(s,5.8,3.4,6.8,2.7,'We inspected the narrow preview, shortened the copy, used a compact paper note and preserved the sea.\nDrift now offers a clear write path and an honest retry. A retry never invents a stranger.',22)
text(s,5.8,6.3,6.8,.45,'Open the actual Make file',17,link=FIGMA)
s.notes_slide.notes_text_frame.text='Figma Make was used for a late-stage iteration after the main game already existed. Version 1 build time reported by Make: 1m32s; revision: 33s. Those are tool-reported generation times, not measured team time savings. The prototype is labeled sample content and has no live game backend. File: '+FIGMA

s=slide('Stage log / 1–2','Choose tools for the question in front of you.',4)
panel(s,.55,2.4,5.95,4.1,'Ideation / ~2 active minutes / 4 of 5','User references + existing code; Figma not used for initial ideation. Scoped this later iteration to empty-pool recovery and returned-story recognition. The rating is Codex’s assessment, not a user-study result.')
panel(s,6.82,2.4,5.95,4.1,'Prototyping / ~4 active minutes / 3 of 5','Figma Make produced four inspectable states from a sanitized brief and existing licensed images. Visual review found an oversized opaque empty state. A second prompt fixed the copy and paper treatment. No fake live activity.')
s.notes_slide.notes_text_frame.text='Stage log covers this late-stage iteration. Minutes are approximate active-time estimates from the session, exclude unattended waiting, and may overlap code work. Each stage is under 200 words. Friction includes a cached preview frame and first-pass layout that needed refinement.'

s=slide('Stage log / 3–4','Trace the design decision into working code.',5)
panel(s,.55,2.4,5.95,4.1,'Handoff / ~3 active minutes / 4 of 5','Selected compact “The water is quiet” feedback and the red ribbon/seal for returned stories. Retained Drift’s existing paper mesh instead of blindly replacing it with the prototype’s timing. Make chat + code diff preserve the path.')
panel(s,6.82,2.4,5.95,4.1,'Code / ~5 active minutes / 4 of 5','Codex implemented React/CSS changes in the existing app; Make code was not copied wholesale. Isolated mobile/desktop checks verify empty catch and retry consume no round, write recovery works and drafts survive reload. Deployed to Cloudflare/Maincloud.')
s.notes_slide.notes_text_frame.text='The estimates describe active time and are not a comparative productivity experiment. Handoff commit: 2544990715d2e2019db9c2f224b35d7e76cd554b. Existing implementation uses a 1.9-second connected mesh opening and immediate reduced-motion fallback; Make suggested 1.1 seconds, which was not adopted without justification. Working code: '+REPO

s=slide('Evidence and limits','Show what works. Name what remains.',6)
panel(s,.55,2.4,5.95,3.95,'Verified software evidence','43 server tests and 18 unit tests pass.\nHosted Claude two-browser flow passes.\nForged Turnstile tokens are rejected.\n360px / 1440px empty-state and draft checks pass.')
panel(s,6.82,2.4,5.95,3.95,'Evidence boundaries','No attendee pilot or social-outcome study.\nNo physical-phone validation yet.\n35/36 synthetic moderation labels matched; hints reviewed by Codex.\nHuman starter inventory remains empty.')
s.notes_slide.notes_text_frame.text='A 50-case real model evaluation produced valid outputs; moderation agreed with 35 of 36 development labels. These are synthetic labels and agent reviews, not independent moderation-quality certification. Source evidence lives in docs/evidence and evals in the private repository.'

s=slide('Working app + traceable files','Open it. Try it. Inspect the handoff.',7)
text(s,.6,2.6,11.9,.6,'Live game',23,bold=True)
text(s,.6,3.15,11.9,.55,APP,17,link=APP)
text(s,.6,4,11.9,.6,'Figma Make interaction prototype · Anyone can view',23,bold=True)
text(s,.6,4.55,11.9,.55,'Drift Bottle Interaction Prototype',19,link=FIGMA)
text(s,.6,5.4,11.9,.5,'Private code · judge access still needs to be arranged',23,bold=True)
text(s,.6,5.95,11.9,.5,REPO,18,link=REPO)
s.notes_slide.notes_text_frame.text='This deck is editable PowerPoint source imported through Figma Slides. No AI imagery was generated. The screenshots are direct browser/test captures. Figma Make and the application are distinct artifacts: one is a labeled interaction prototype, one has the real Spacetime/Claude backend. Mana Bottle Kitchen is a separate content-preparation workflow, not the game backend.'

r.save(OUT)
print(OUT)
